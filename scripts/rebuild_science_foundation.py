#!/usr/bin/env python3
"""Rebuild the NGSS three-dimensional foundation (seps / dcis / cccs) on every PE in
data/science.json from machine-readable NGSS sources.

Why: the foundation originally came from a `pdftotext -layout` pass over the NJSLS-Science
PDFs, which mangled the multi-column foundation tables (intros and bullets split mid-sentence,
dimensions bled together, 15 PEs with empty arrays). Those PDFs are gone from ../curriculum/,
so this re-derives every PE deterministically instead of patching the extraction.

Sources — downloaded from a pinned commit of github.com/galacticpolymath/standardX and
verified by SHA-256. They are deliberately NOT vendored: that repo is GPLv3 and compiles
third-party standards.
  * CodesForNGSSelements_TedWillard.xlsx — the NSTA Atlas of the Three Dimensions element
    coding: PE -> element-code mapping, verbatim element text, canonical DCI names.
  * NGSS.xml — the ASN RDF export of NGSS (Achieve, 2013): the SEP grade-band progression
    intros, the official connection sub-concept names, and the text of the one element the
    workbook omits (INFLU-M4).

Only seps / dcis / cccs are replaced; every other field is left byte-identical. Conventions
match the lens keys build_all.py emits:
  * seps — the 8 practices only; intro = the practice's grade-band progression.
  * dcis — core then secondary core ideas; secondary ones carry "secondary": true.
  * cccs — core crosscutting concepts, then Connections to Engineering, Technology, and
           Applications of Science, then Connections to Nature of Science. `concept` is always
           the specific concept / sub-concept; connections also carry `connection`:
           "engineering" | "nature_of_science". `intro` is always "".

Requires openpyxl (`pip install openpyxl`). Run, then rebuild the flat index:
    python3 scripts/rebuild_science_foundation.py && python3 scripts/build_all.py
"""
import hashlib
import json
import re
import sys
import tempfile
import urllib.request
import xml.etree.ElementTree as ET
from collections import OrderedDict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SCIENCE = ROOT / "data" / "science.json"

UPSTREAM = "https://raw.githubusercontent.com/galacticpolymath/standardX/cbf6666109ffedb564692d796cc447711de4fe0c"
SOURCES = {
    "willard": ("data/CodesForNGSSelements_TedWillard.xlsx",
                "c31fc62108d7c42182f8fb1302af5788a6b889e73737346c67e4b5431f1b631b"),
    "asn":     ("data/NGSS.xml",
                "f0bd4a55de426b54bcd17332454d05b8992ac534f6bf4946f90666bae385695d"),
}
CACHE = Path(tempfile.gettempdir()) / "njsls_ngss_sources"

SEP_NAME = {
    "AQDP": "Asking Questions and Defining Problems",
    "MOD": "Developing and Using Models",
    "INV": "Planning and Carrying Out Investigations",
    "DATA": "Analyzing and Interpreting Data",
    "MATH": "Using Mathematics and Computational Thinking",
    "CEDS": "Constructing Explanations and Designing Solutions",
    "ARG": "Engaging in Argument from Evidence",
    "INFO": "Obtaining, Evaluating, and Communicating Information",
}
CCC_NAME = {
    "PAT": "Patterns", "CE": "Cause and Effect", "SPQ": "Scale, Proportion, and Quantity",
    "SYS": "Systems and System Models", "EM": "Energy and Matter",
    "SF": "Structure and Function", "SC": "Stability and Change",
}
# Official NGSS names (Appendix J / H), as they appear in the ASN export and in the
# NJSLS-derived data — the workbook's own labels vary in wording/casing.
ENG_NAME = {
    "INFLU": "Influence of Engineering, Technology, and Science on Society and the Natural World",
    "INTER": "Interdependence of Science, Engineering, and Technology",
}
NOS_NAME = {
    "VOM": "Scientific Investigations Use a Variety of Methods",
    "BEE": "Scientific Knowledge is Based on Empirical Evidence",
    "OTR": "Scientific Knowledge is Open to Revision in Light of New Evidence",
    "ENP": "Science Models, Laws, Mechanisms, and Theories Explain Natural Phenomena",
    "WOK": "Science is a Way of Knowing",
    "AOC": "Scientific Knowledge Assumes an Order and Consistency in Natural Systems",
    "HE": "Science is a Human Endeavor",
    "AQAW": "Science Addresses Questions About the Natural and Material World",
}
# The workbook's Eng Codes sheet stops at INFLU-M3; MS-PS4-3 cites INFLU-M4.
INFLU_M4 = ("Technologies extend the measurement, exploration, modeling, and computational "
            "capacity of scientific investigations.")

CODE_RE = re.compile(r"[A-Z][A-Z0-9.]*-[PEMH]\d+")
WARN = []


def fetch(key):
    rel, sha = SOURCES[key]
    path = CACHE / Path(rel).name
    if not path.exists() or hashlib.sha256(path.read_bytes()).hexdigest() != sha:
        CACHE.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(f"{UPSTREAM}/{rel}", timeout=60) as r:
            path.write_bytes(r.read())
    got = hashlib.sha256(path.read_bytes()).hexdigest()
    if got != sha:
        sys.exit(f"SHA-256 mismatch for {rel}: expected {sha}, got {got}")
    return path


def cell(x):
    return str(x).strip() if x is not None else ""


def codes(raw):
    """Element codes in a workbook cell. Repairs the two source typos first: a comma used for
    the period inside a DCI code ('PS1,A-M5', MS-PS1-1); a period used for the comma between
    codes ('LS3.A-M1. LS3.B-M2', MS-LS3-1) is handled by extracting with a regex."""
    return CODE_RE.findall(re.sub(r"([A-Z]{2,4}\d),(?=[A-Z])", r"\1.", cell(raw)))


def load_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = lambda name: [list(r) for r in wb[name].iter_rows(values_only=True)]
    text = {}
    for sheet in ("SEP Codes", "DCI Codes", "CCC Codes", "Eng Codes", "NOS Codes"):
        for r in rows(sheet)[1:]:
            if r and cell(r[0]) and cell(r[1]):
                text[cell(r[0])] = cell(r[1])
    dci_name = {}
    for r in rows("Topic Codes"):
        m = re.match(r"^([A-Z]{2,4}\d\.[A-Z]):?$", cell(r[0]) if r else "")
        if m and len(r) > 1 and cell(r[1]):
            dci_name[m.group(1)] = cell(r[1])
    pe_map = {}
    for r in rows("PE Foundations and Connections")[1:]:
        if cell(r[1]):
            pe_map[cell(r[1])] = {k: codes(r[i]) for k, i in
                                  (("practices", 6), ("core", 7), ("secondary", 8),
                                   ("ccc", 9), ("eng", 10), ("nos", 11))}
    return text, dci_name, pe_map


def load_asn(path):
    ns = {"asn": "http://purl.org/ASN/schema/core/", "dcterms": "http://purl.org/dc/terms/",
          "gemq": "http://purl.org/gem/qualifiers/", "rdf": "http://www.w3.org/1999/02/22-rdf-syntax-ns#"}
    rdf_res = "{%s}resource" % ns["rdf"]
    nodes = {}
    for st in ET.parse(path).getroot().findall("asn:Statement", ns):
        desc = st.find("dcterms:description", ns)
        parent = st.find("gemq:isChildOf", ns)
        nodes[st.get("{%s}about" % ns["rdf"])] = {
            "desc": desc.text.strip() if desc is not None and desc.text else "",
            "parent": parent.get(rdf_res) if parent is not None else None,
            "edu": {e.get(rdf_res).rsplit("/", 1)[-1] for e in st.findall("dcterms:educationLevel", ns)},
        }
    intros = {}
    for n in nodes.values():
        if "progresses to" not in n["desc"]:
            continue
        band = "E" if {"3", "4", "5"} <= n["edu"] else "M" if {"6", "7", "8"} <= n["edu"] else None
        practice = (nodes.get(n["parent"]) or {}).get("desc")
        if band and practice:
            intros[(practice, band)] = n["desc"]
    # The 3-5 Using-Mathematics progression lacks a clean 3-5 educationLevel tag.
    for n in nodes.values():
        if n["desc"].startswith("Mathematical and computational thinking") and "3–5" in n["desc"]:
            intros.setdefault(("Using Mathematics and Computational Thinking", "E"), n["desc"])
    all_desc = {n["desc"] for n in nodes.values()}
    for name in (*ENG_NAME.values(), *(v for k, v in NOS_NAME.items() if k != "WOK")):
        if name not in all_desc:
            WARN.append(f"connection name not found verbatim in ASN: {name}")
    if not any(INFLU_M4 in d for d in all_desc):
        WARN.append("INFLU-M4 text not found verbatim in ASN")
    return intros


def grouped(element_codes, names):
    """Ordered {prefix: [codes]} for codes whose prefix is in `names`, numeric order within."""
    out = OrderedDict()
    for c in element_codes:
        p = c.rsplit("-", 1)[0]
        if p in names:
            out.setdefault(p, []).append(c)
    for cs in out.values():
        cs.sort(key=lambda c: int(re.search(r"(\d+)$", c).group(1)))
    return out


def rebuild_pe(code, m, text, dci_name, intros):
    def t(c):
        if c not in text:
            WARN.append(f"{code}: no text for element {c}")
        return text.get(c, c)

    seps = []
    for p, cs in grouped(m["practices"], SEP_NAME).items():  # drops stray non-SEP codes (MS-ESS1-4)
        band = re.search(r"-([PEMH])\d+$", cs[0]).group(1)
        intro = intros.get((SEP_NAME[p], band), "")
        if not intro:
            WARN.append(f"{code}: no progression intro for {SEP_NAME[p]} ({band})")
        seps.append({"practice": SEP_NAME[p], "intro": intro, "bullets": [t(c) for c in cs]})

    dcis = []
    for key, secondary in (("core", False), ("secondary", True)):
        for subcat, cs in grouped(m[key], dci_name).items():
            entry = {"code": subcat, "name": dci_name[subcat], "bullets": [t(c) for c in cs]}
            if secondary:
                entry["secondary"] = True
            dcis.append(entry)

    cccs = [{"concept": CCC_NAME[p], "intro": "", "bullets": [t(c) for c in cs]}
            for p, cs in grouped(m["ccc"], CCC_NAME).items()]
    for codes_, names, kind in ((m["eng"], ENG_NAME, "engineering"),
                                (m["nos"], NOS_NAME, "nature_of_science")):
        for p, cs in grouped(codes_, names).items():
            cccs.append({"concept": names[p], "intro": "", "connection": kind,
                         "bullets": [t(c) for c in cs]})
    return seps, dcis, cccs


def main():
    text, dci_name, pe_map = load_workbook(fetch("willard"))
    text.setdefault("INFLU-M4", INFLU_M4)
    intros = load_asn(fetch("asn"))

    sci = json.loads(SCIENCE.read_text())
    rebuilt, regressed, empty = 0, [], []
    for disc in sci.values():
        for topic in disc["topics"]:
            for pe in topic["pes"]:
                m = pe_map.get(pe["code"])
                if not m:
                    WARN.append(f"{pe['code']}: not in the workbook's PE map — left as-is")
                    continue
                had = {k: bool(pe.get(k)) for k in ("seps", "dcis", "cccs")}
                pe["seps"], pe["dcis"], pe["cccs"] = rebuild_pe(pe["code"], m, text, dci_name, intros)
                for k in ("seps", "dcis", "cccs"):
                    if not pe[k]:
                        (regressed if had[k] else empty).append(f"{pe['code']}.{k}")
                rebuilt += 1

    SCIENCE.write_text(json.dumps(sci, indent=2, ensure_ascii=False) + "\n")
    print(f"Rebuilt the foundation on {rebuilt} PEs -> {SCIENCE.relative_to(ROOT)}")
    print(f"Empty dimensions (NGSS lists none — the ETS1 design PEs have no CCC): {empty or 'none'}")
    if regressed:
        print(f"!! Previously-populated dimensions now empty: {regressed}")
    for w in WARN:
        print(f"!! {w}")
    sys.exit(1 if regressed or WARN else 0)


if __name__ == "__main__":
    main()
